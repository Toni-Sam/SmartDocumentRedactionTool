"""Compare original vs redacted documents and score the redaction.

Usage (from the folder containing the files):
    python evaluate_redaction.py [folder_with_files]

Metrics
  * Item recall   - ground-truth PII items no longer present in the output text.
  * Word-level P/R/F1 - words removed by the tool vs words that are PII
    (precision drops when non-PII words are removed, recall drops when PII leaks).
Ground truth is written by hand below. Edit the lists if you change the test files.
"""
import re, json, sys, subprocess
from collections import Counter
from pathlib import Path

try:
    import fitz  # PyMuPDF
except ImportError:
    fitz = None
from docx import Document

DATA = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(".")  # folder with the test files


def docx_text(path):
    d = Document(path)
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for r in t.rows:
            for c in r.cells:
                parts.append(c.text)
    return "\n".join(parts)


def pdf_text(path):
    if fitz:
        with fitz.open(path) as doc:
            return "\n".join(p.get_text() for p in doc)
    return subprocess.run(["pdftotext", str(path), "-"], capture_output=True, text=True).stdout


def norm(s):
    return re.sub(r"\s+", " ", s).strip()


def words(s):
    toks = (w.strip(".,;:()\"'") for w in norm(s).split(" "))
    return [w for w in toks if w and w != "[REDACTED]"]


DOCX_PII = {
    "Person name": ["Omitogun Oluwatoni", "Oladele Peter", "Sanni Ayomide", "Shittu Mariam"],
    "NIN": ["49340714800", "49340714800", "45612349219", "41171936790"],
    "BVN": ["45189023173", "22331019874", "22154671896", "10139492088"],
    "Voter no.": ["804R DFT9 2617 3628 1038", "361E QTR4 9172 1628 2618", "152F QIA1 2819 6291 1028"],
    "SSN (US)": ["225-44-1989"],
    "Address": ["2, Akeem Shittu Street", "Alao-Akala Way", "Akobo", "Ibadan", "Oyo", "Nigeria"],
}
PDF_PII = {
    "Person name": ["Adaeze Chinonso Okonkwo", "Adaeze Okonkwo", "Ifeoma Balogun"],
    "NIN": ["49340714800", "49340714800", "58273649102"],
    "BVN": ["22157896634", "22157896634", "30987654321"],
    "Voter no.": ["90AB CD12 3456 7890 XY12", "90AB CD12 3456 7890 XY12", "71CD EF34 5678 9012 AB34"],
    "SSN (US)": ["225-44-1989"],
    "Address": ["15 Herbert Macaulay Way", "Yaba", "Lagos"],
    "State / LGA": ["Anambra", "Awka South"],
    "Email": ["adaeze.okonkwo@example.com"],
    "NUBAN": ["0123456789"],
    "Phone": ["08031234567"],
    "Passport": ["A12345678"],
    "Driver licence": ["ABC12345DE67"],
    "Plate": ["ABJ-123KJ"],
    "TIN": ["12345678-0001"],
    "RSA PIN": ["PEN123456789012"],
    "NYSC call-up": ["NYSC/LAG/2024/123456"],
    "NYSC state code": ["OY/24A/1234"],
    "CAC RC": ["RC1234567"],
    "NHIS": ["0987654321"],
}


def present(item, text):
    return re.search(r"(?<![\w-])" + re.escape(norm(item)) + r"(?![\w-])", norm(text)) is not None


def score(orig, red, pii):
    per_type, leaked = {}, []
    for k, items in pii.items():
        hit = [i for i in items if present(i, red)]
        per_type[k] = (len(items) - len(hit), len(items))
        leaked += hit
    pii_words = Counter(w for items in pii.values() for i in items for w in words(i))
    removed = Counter(words(orig)) - Counter(words(red))
    tp = sum((removed & pii_words).values())
    fp_words = removed - pii_words
    fn = sum((pii_words - removed).values())
    fp = sum(fp_words.values())
    p = tp / (tp + fp) if tp + fp else 0
    r = tp / (tp + fn) if tp + fn else 0
    f1 = 2 * p * r / (p + r) if p + r else 0
    ok = sum(v[0] for v in per_type.values()); total = sum(v[1] for v in per_type.values())
    return dict(items_ok=ok, items_total=total, leaked=leaked, per_type=per_type,
                tp=tp, fp=fp, fn=fn, precision=p, recall=r, f1=f1,
                collateral=sorted(fp_words.elements()))


RUNS = {
    "DOCX \u2192 DOCX": (docx_text(DATA / "DOCX_Redaction_Test.docx"), docx_text(DATA / "DOCX_Redaction_Test_redacted.docx"), DOCX_PII),
    "DOCX \u2192 PDF": (docx_text(DATA / "DOCX_Redaction_Test.docx"), pdf_text(DATA / "DOCX_Redaction_Test_Redacted.pdf"), DOCX_PII),
    "PDF \u2192 PDF": (pdf_text(DATA / "PDF_Redaction_Test.pdf"), pdf_text(DATA / "PDF_Redaction_Test_redacted.pdf"), PDF_PII),
}

if __name__ == "__main__":
    results = {n: score(*a) for n, a in RUNS.items()}
    for n, r in results.items():
        print(f"\n== {n}")
        print(f"  PII items fully removed: {r['items_ok']}/{r['items_total']}   leaked: {r['leaked']}")
        print(f"  word-level  P={r['precision']:.2f}  R={r['recall']:.2f}  F1={r['f1']:.2f}  (TP {r['tp']}, FP {r['fp']}, FN {r['fn']})")
        print(f"  non-PII words lost ({r['fp']}): {r['collateral']}")
    json.dump(results, open("results.json", "w"), indent=1, ensure_ascii=False)